"""Tests for KimiAgent.

The emphasis is on the properties that actually break decks and are invisible
in a smoke test:

* **Nothing overflows.** Text is measured before it is drawn, so every text box
  in a rendered ``.pptx`` can be re-measured and asserted to fit.
* **Nothing is silently dropped.** Pagination must preserve every item.
* **Nothing is unreadable.** Every theme's colour pairings must clear WCAG.
* **Old specs still work.** The v1 spec format must keep rendering.
"""

from __future__ import annotations

import json
import os

import pytest

from kimiagent import design as D
from kimiagent.agent import DeckAgent, critique
from kimiagent.layout import (
    MAX_AGENDA,
    card_grid,
    paginate,
    stat_grid,
    zones,
)
from kimiagent.models import (
    Bullet,
    Card,
    ChartData,
    ChartSeries,
    ComparisonSide,
    Deck,
    ProcessStep,
    Slide,
    SlideType,
    Stat,
    TableData,
    TimelineItem,
)
from kimiagent.pdf_renderer import render_pdf
from kimiagent.renderer import render_deck
from kimiagent.themes import THEMES, Theme, get_theme, theme_names

pptx = pytest.importorskip("pptx")
EMU_PER_PT = 12700


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------


def _full_deck(theme: str = "kimi", **kw) -> Deck:
    """A deck touching every archetype."""
    return Deck(
        title="Every Archetype",
        subtitle="A test deck",
        author="Tests",
        theme=theme,
        footer_note="Internal",
        slides=[
            Slide(type=SlideType.COVER, title="Every Archetype", subtitle="A test deck", kicker="Test"),
            Slide(type=SlideType.AGENDA, title="Agenda",
                  bullets=[Bullet(text=f"Item {i}") for i in range(1, 6)]),
            Slide(type=SlideType.SECTION, title="First section", subtitle="01"),
            Slide(type=SlideType.BIG_STATEMENT, title="One claim worth a slide.",
                  kicker="Thesis", subtitle="Supporting line"),
            Slide(type=SlideType.STATS, title="Numbers", kicker="At a glance", stats=[
                Stat(value="-71%", label="Cost", detail="since 2015", trend="down"),
                Stat(value="4.2x", label="Deployed", detail="vs 2021", trend="up"),
                Stat(value="38%", label="Share", trend="flat"),
            ]),
            Slide(type=SlideType.CHART, title="Trend", chart=ChartData(
                chart_type="line", categories=["2024", "2025", "2026"],
                series=[ChartSeries("A", [1, 2, 3]), ChartSeries("B", [3, 2, 1])],
                takeaway="A overtook B.")),
            Slide(type=SlideType.CHART, title="Split", chart=ChartData(
                chart_type="doughnut", data_labels=True, categories=["x", "y", "z"],
                series=[ChartSeries("S", [50, 30, 20])])),
            Slide(type=SlideType.CARDS, title="Cards", cards=[
                Card(title="One", body="First idea.", icon="1"),
                Card(title="Two", body="Second idea.", icon="2"),
                Card(title="Three", body="Third idea."),
            ]),
            Slide(type=SlideType.TIMELINE, title="Timeline", timeline=[
                TimelineItem(period="2024", title="Start", detail="Kickoff", done=True),
                TimelineItem(period="2025", title="Middle", detail="Scale"),
                TimelineItem(period="2026", title="End", detail="Payoff"),
            ]),
            Slide(type=SlideType.PROCESS, title="Process", steps=[
                ProcessStep(title="Interconnect", detail="Queue position."),
                ProcessStep(title="Finance", detail="Debt sizing."),
                ProcessStep(title="Build", detail="Construction."),
            ]),
            Slide(type=SlideType.COMPARISON, title="Choice", verdict="Pick B.", sides=[
                ComparisonSide(title="A", subtitle="Conventional",
                               points=[Bullet(text="Fast"), Bullet(text="Limited", bold_lead="Risk")]),
                ComparisonSide(title="B", subtitle="Better", preferred=True,
                               points=[Bullet(text="Compounds")]),
            ]),
            Slide(type=SlideType.TABLE, title="Table", table=TableData(
                headers=["Region", "GW", "Rate"], numeric_columns=[1, 2], highlight_row=2,
                rows=[["PJM", "286", "14%"], ["MISO", "241", "11%"], ["Total", "527", "13%"]])),
            Slide(type=SlideType.BULLETS, title="Bullets", subtitle="With nesting", bullets=[
                Bullet(text="Top level with a lead-in.", bold_lead="Lead"),
                Bullet(text="Plain top level point"),
                Bullet(text="A nested supporting detail", level=1),
            ], notes="Speaker notes here."),
            Slide(type=SlideType.TWO_COLUMN, title="Two columns", columns=[
                [Bullet(text="Left one"), Bullet(text="Left two")],
                [Bullet(text="Right one"), Bullet(text="Right two")],
            ]),
            Slide(type=SlideType.QUOTE, quote="A memorable line about the work.",
                  attribution="Someone credible"),
            Slide(type=SlideType.IMAGE, title="Image", caption="Figure 1 — a diagram"),
            Slide(type=SlideType.CLOSING, title="Thank you", subtitle="Questions"),
        ],
        **kw,
    )


# --------------------------------------------------------------------------
# design.py
# --------------------------------------------------------------------------


class TestDesignMetrics:
    def test_width_matches_real_helvetica(self):
        """Embedded metrics must match the font the PDF backend actually uses."""
        from reportlab.pdfbase.pdfmetrics import stringWidth

        for text in ("Hello Modern Deck", "iiiiWWWW", "The State of Renewable Energy"):
            for bold in (False, True):
                font = "Helvetica-Bold" if bold else "Helvetica"
                assert D.text_width(text, 24, bold) == pytest.approx(
                    stringWidth(text, font, 24), abs=0.01
                )

    def test_wrap_respects_width(self):
        long = "Distributed systems fail in correlated ways, not independent ones."
        for width in (120, 240, 480):
            for line in D.wrap_text(long, width, 14):
                assert D.text_width(line, 14) <= width + 0.01

    def test_wrap_hard_breaks_unbreakable_words(self):
        lines = D.wrap_text("A" * 200, 100, 14)
        assert len(lines) > 1
        assert all(D.text_width(l, 14) <= 100.01 for l in lines)

    def test_fit_font_size_fits(self):
        long = "Distributed systems fail in correlated ways, " * 3
        size = D.fit_font_size(long, 400, 90, ideal=24)
        assert size <= 24
        assert D.text_height(long, 400, size) <= 90

    def test_fit_avoids_mid_word_breaks(self):
        """A long word in a narrow box must shrink, not split."""
        size = D.fit_font_size("Interconnection", 70, 200, ideal=19)
        assert D.longest_word_width("Interconnection", size) <= 70
        assert len(D.wrap_text("Interconnection", 70, size)) == 1

    def test_truncate_only_when_needed(self):
        assert D.truncate_to_fit("short", 300, 100, 12) == "short"
        out = D.truncate_to_fit("word " * 100, 200, 40, 12)
        assert D.text_height(out, 200, 12) <= 40

    def test_cjk_is_full_width(self):
        assert D.text_width("渲染现代化", 20) == pytest.approx(100.0)


class TestDesignColour:
    def test_contrast_ratio_bounds(self):
        assert D.contrast_ratio("FFFFFF", "000000") == pytest.approx(21.0, abs=0.01)
        assert D.contrast_ratio("AABBCC", "AABBCC") == pytest.approx(1.0, abs=0.01)

    def test_readable_text_picks_legible_side(self):
        assert D.readable_text_on("FFD166") != "FFFFFF"   # pale amber needs dark ink
        assert D.readable_text_on("1B2A4A") == "FFFFFF"   # deep navy needs light

    def test_ensure_contrast_reaches_threshold(self):
        fixed = D.ensure_contrast("999999", "FFFFFF", 4.5)
        assert D.contrast_ratio(fixed, "FFFFFF") >= 4.5

    def test_grid_arithmetic_is_exact(self):
        g = D.GRID
        assert 2 * g.span(6) + g.gutter == pytest.approx(g.content_w)
        assert g.col(0) == pytest.approx(g.left)
        row = g.row(4)
        assert row[-1][0] + row[-1][1] == pytest.approx(g.right)


# --------------------------------------------------------------------------
# themes.py
# --------------------------------------------------------------------------


class TestThemes:
    @pytest.mark.parametrize("name", theme_names())
    def test_body_text_is_readable(self, name):
        t = THEMES[name]
        assert D.contrast_ratio(t.text, t.background) >= 7.0
        assert D.contrast_ratio(t.text, t.surface) >= 4.5

    @pytest.mark.parametrize("name", theme_names())
    def test_hero_gradient_is_text_safe(self, name):
        """One text colour must be legible across the whole gradient sweep."""
        t = THEMES[name]
        start, end, _ = t.gradient
        ink = t.hero_text
        worst = min(
            D.contrast_ratio(ink, start),
            D.contrast_ratio(ink, end),
            D.contrast_ratio(ink, D.mix(start, end, 0.5)),
        )
        assert worst >= 4.5

    @pytest.mark.parametrize("name", theme_names())
    def test_large_accent_text_is_readable(self, name):
        t = THEMES[name]
        assert D.contrast_ratio(t.accent_on(t.background), t.background) >= 3.0
        assert D.contrast_ratio(t.on(t.primary), t.primary) >= 4.5

    def test_legacy_theme_only_original_fields(self):
        """A theme predating the new tokens must still construct and render."""
        legacy = Theme(
            name="legacy", primary="123456", secondary="ABCDEF", background="FFFFFF",
            surface="EEEEEE", text="000000", text_muted="777777", text_on_primary="FFFFFF",
        )
        assert legacy.border and legacy.chart_grid and legacy.surface_alt
        assert legacy.gradient_from and legacy.gradient_to
        assert legacy.accents == ["ABCDEF"]
        assert legacy.dark is False

    def test_pathological_gradient_still_resolves(self):
        """An all-light gradient must be made legible rather than left broken."""
        t = Theme(
            name="x", primary="FFFF00", secondary="FFFFE0", background="FFFFFF",
            surface="FFFFF0", text="222200", text_muted="888800", text_on_primary="000000",
            gradient_from="FFFF00", gradient_to="FFFFF0",
        )
        start, end, _ = t.gradient
        assert min(D.contrast_ratio(t.hero_text, start),
                   D.contrast_ratio(t.hero_text, end)) >= 4.5

    def test_variant_does_not_mutate_original(self):
        t = get_theme("slate")
        assert t.variant(radius=0).radius == 0
        assert t.radius != 0

    def test_single_series_gets_monochrome_ramp(self):
        t = get_theme("slate")
        assert len(set(t.series_colors(4, monochrome=True))) == 4
        assert t.series_colors(3) == [t.accent(0), t.accent(1), t.accent(2)]

    def test_unknown_theme_falls_back(self):
        assert get_theme("nope").name == "kimi"
        assert get_theme("").name == "kimi"


# --------------------------------------------------------------------------
# models.py
# --------------------------------------------------------------------------


class TestModels:
    def test_round_trip_is_stable(self):
        deck = _full_deck()
        once = Deck.from_dict(deck.to_dict())
        twice = Deck.from_dict(json.loads(json.dumps(once.to_dict())))
        assert once.to_dict() == twice.to_dict()

    def test_v1_spec_still_loads(self):
        """The spec format documented in the first release must keep working."""
        legacy = {
            "title": "My Deck", "theme": "kimi",
            "slides": [
                {"type": "cover", "title": "My Deck", "subtitle": "A subtitle"},
                {"type": "bullets", "title": "Key points", "bullets": ["First", "Second"]},
                {"type": "stats", "title": "Impact",
                 "stats": [{"value": "42%", "label": "growth", "detail": "YoY"}]},
                {"type": "chart", "title": "Trend", "chart": {
                    "chart_type": "line", "categories": ["2024", "2025"],
                    "series": [{"name": "Revenue", "values": [10, 18]}]}},
                {"type": "closing", "title": "Thank you"},
            ],
        }
        deck = Deck.from_dict(legacy)
        assert [s.type.value for s in deck.slides] == [
            "cover", "bullets", "stats", "chart", "closing"
        ]
        assert [b.text for b in deck.slides[1].bullets] == ["First", "Second"]

    @pytest.mark.parametrize("alias,expected", [
        ("Title Slide", "cover"), ("pros-cons", "comparison"), ("KPIs", "stats"),
        ("roadmap", "timeline"), ("feature_cards", "cards"), ("workflow", "process"),
        ("graph", "chart"), ("thank_you", "closing"), ("totally_unknown", "bullets"),
    ])
    def test_type_aliases(self, alias, expected):
        assert Slide.from_dict({"type": alias}).type.value == expected

    def test_coerces_messy_llm_output(self):
        deck = Deck.from_dict({"title": "M", "slides": [
            {"type": "stats", "stats": [{"value": 42, "label": "n", "trend": "UP"}]},
            {"type": "chart", "chart": {"chart_type": "column", "categories": [2024, 2025],
                                        "series": [{"name": "S", "values": ["1,240", "18%", None]}]}},
            {"type": "table", "table": {"headers": ["a"], "rows": [[1], [None]]}},
            {"type": "cards", "cards": ["a bare string"]},
        ]})
        assert deck.slides[0].stats[0].value == "42"
        assert deck.slides[0].stats[0].trend == "up"
        assert deck.slides[1].chart.series[0].values == [1240.0, 18.0, 0.0]
        assert deck.slides[1].chart.categories == ["2024", "2025"]
        assert deck.slides[2].table.rows == [["1"], [""]]
        assert deck.slides[3].cards[0].title == "a bare string"

    def test_missing_and_null_slides(self):
        assert Deck.from_dict({"title": "a"}).slide_count() == 0
        assert Deck.from_dict({"title": "a", "slides": None}).slide_count() == 0

    def test_variety_score(self):
        all_bullets = Deck(title="x", slides=[Slide(type=SlideType.BULLETS) for _ in range(4)])
        assert all_bullets.variety_score() == 0.0
        assert _full_deck().variety_score() > 0.7


# --------------------------------------------------------------------------
# layout.py — pagination must never lose content
# --------------------------------------------------------------------------


class TestPagination:
    def _counts(self, deck):
        return {
            "bullets": sum(len(s.bullets) for s in deck.slides),
            "stats": sum(len(s.stats) for s in deck.slides),
            "cards": sum(len(s.cards) for s in deck.slides),
            "timeline": sum(len(s.timeline) for s in deck.slides),
            "steps": sum(len(s.steps) for s in deck.slides),
            "rows": sum(len(s.table.rows) for s in deck.slides if s.table),
            "columns": sum(len(c) for s in deck.slides for c in s.columns),
            "points": sum(len(x.points) for s in deck.slides for x in s.sides),
        }

    @pytest.fixture
    def overflowing(self):
        long = "Distributed systems fail in correlated ways, which is why redundancy alone never buys the availability you expect."
        return Deck(title="Overflow", slides=[
            Slide(type=SlideType.BULLETS, title="Many", notes="N",
                  bullets=[Bullet(text=long) for _ in range(10)]),
            Slide(type=SlideType.AGENDA, title="Agenda",
                  bullets=[Bullet(text=f"i{i}") for i in range(15)]),
            Slide(type=SlideType.STATS, title="S",
                  stats=[Stat(value=str(i), label="l") for i in range(11)]),
            Slide(type=SlideType.CARDS, title="C",
                  cards=[Card(title=f"c{i}") for i in range(14)]),
            Slide(type=SlideType.TIMELINE, title="T",
                  timeline=[TimelineItem(period=str(i), title=f"t{i}") for i in range(12)]),
            Slide(type=SlideType.PROCESS, title="P",
                  steps=[ProcessStep(title=f"s{i}") for i in range(11)]),
            Slide(type=SlideType.TABLE, title="Tb", table=TableData(
                headers=["a", "b"], highlight_row=20,
                rows=[[str(i), "x"] for i in range(25)])),
            Slide(type=SlideType.TWO_COLUMN, title="2c", columns=[
                [Bullet(text=f"l{i}") for i in range(13)], [Bullet(text="r")]]),
            Slide(type=SlideType.COMPARISON, title="Cmp", verdict="V", sides=[
                ComparisonSide(title="A", points=[Bullet(text=f"a{i}") for i in range(14)]),
                ComparisonSide(title="B", points=[Bullet(text="b")])]),
        ])

    def test_nothing_is_lost(self, overflowing):
        before = self._counts(overflowing)
        after = self._counts(paginate(overflowing))
        assert before == after

    def test_it_actually_splits(self, overflowing):
        assert paginate(overflowing).slide_count() > overflowing.slide_count()

    def test_original_deck_is_not_mutated(self, overflowing):
        n = overflowing.slide_count()
        paginate(overflowing)
        assert overflowing.slide_count() == n

    def test_continuation_titles_do_not_stack(self, overflowing):
        for s in paginate(overflowing).slides:
            assert s.title.count("(cont.)") <= 1

    def test_notes_and_verdict_only_on_first_page(self, overflowing):
        pages = [s for s in paginate(overflowing).slides if s.title.startswith("Cmp")]
        assert pages[0].verdict == "V"
        assert all(not p.verdict for p in pages[1:])

    def test_table_highlight_row_is_remapped(self, overflowing):
        for s in paginate(overflowing).slides:
            if s.table and s.table.highlight_row is not None:
                assert s.table.rows[s.table.highlight_row] == ["20", "x"]

    def test_agenda_respects_cap(self, overflowing):
        for s in paginate(overflowing).slides:
            if s.type is SlideType.AGENDA:
                assert len(s.bullets) <= MAX_AGENDA

    def test_bullet_capacity_is_content_aware(self):
        body = zones().body
        from kimiagent.layout import bullet_capacity

        short = [Bullet(text="Short") for _ in range(12)]
        long = [Bullet(text="A considerably longer bullet that is guaranteed to wrap onto "
                            "several lines because it keeps going well past any sensible "
                            "length for a slide, continuing at length about nothing in "
                            "particular, and then continuing further still.") for _ in range(12)]
        assert bullet_capacity(short, body) > bullet_capacity(long, body)


class TestZonesAndGrids:
    @pytest.mark.parametrize("kicker", [False, True])
    @pytest.mark.parametrize("subtitle", [False, True])
    @pytest.mark.parametrize("lines", [1, 2])
    def test_zones_never_overlap(self, kicker, subtitle, lines):
        z = zones(has_kicker=kicker, has_subtitle=subtitle, title_lines=lines)
        assert z.title.y >= z.kicker.bottom - 0.01
        assert z.rule.y >= z.title.bottom - 0.01
        assert z.body.y >= z.rule.bottom
        assert z.body.bottom <= z.footer.y + 0.01
        assert z.body.h > 100

    @pytest.mark.parametrize("n", range(1, 10))
    def test_card_grid_stays_in_bounds(self, n):
        body = zones().body
        rects = card_grid(body, n)
        assert len(rects) == n
        for r in rects:
            assert r.x >= body.x - 0.01 and r.right <= body.right + 0.01
            assert r.y >= body.y - 0.01 and r.bottom <= body.bottom + 0.01
        # Uniform card width reads as a grid; mixed widths read as a mistake.
        assert len({round(r.w, 1) for r in rects}) == 1

    def test_card_grid_balances_four_as_two_by_two(self):
        rects = card_grid(zones().body, 4)
        assert len({round(r.y, 1) for r in rects}) == 2

    @pytest.mark.parametrize("n", range(1, 7))
    def test_stat_grid_in_bounds(self, n):
        body = zones().body
        for r in stat_grid(body, n):
            assert r.x >= body.x - 0.01 and r.right <= body.right + 0.01


# --------------------------------------------------------------------------
# Renderers
# --------------------------------------------------------------------------


def _measure_pptx_overflow(path):
    """Re-measure every text box in a saved deck. Returns a list of overflows."""
    from pptx import Presentation

    prs = Presentation(path)
    slide_w = prs.slide_width / EMU_PER_PT
    slide_h = prs.slide_height / EMU_PER_PT
    overflow, out_of_bounds = [], []

    for idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            w = shape.width / EMU_PER_PT
            h = shape.height / EMU_PER_PT
            x = shape.left / EMU_PER_PT
            y = shape.top / EMU_PER_PT
            if x < -1 or y < -1 or x + w > slide_w + 1 or y + h > slide_h + 1:
                out_of_bounds.append((idx, shape.text_frame.text[:40]))

            total = 0.0
            for p in shape.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not text:
                    continue
                size, bold = None, False
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                    bold = bold or bool(r.font.bold)
                size = size or 18
                spacing = p.line_spacing if isinstance(p.line_spacing, (int, float)) else 1.3
                margin = 0.0
                pPr = p._pPr
                if pPr is not None and pPr.get("marL"):
                    margin = int(pPr.get("marL")) / EMU_PER_PT
                total += D.text_height(text, max(1.0, w - margin), size, bold, False, spacing)
                total += p.space_after.pt if p.space_after else 0
            if total > h + 1.5:
                overflow.append((idx, round(total - h, 1), shape.text_frame.text[:40]))
    return overflow, out_of_bounds


class TestPptxRenderer:
    @pytest.mark.parametrize("theme", theme_names())
    def test_no_text_overflows_any_theme(self, theme, tmp_path):
        """The core guarantee: measured text always fits the box it was given."""
        out = render_deck(_full_deck(theme), str(tmp_path / f"{theme}.pptx"))
        overflow, oob = _measure_pptx_overflow(out)
        assert not overflow, f"overflowing text: {overflow[:4]}"
        assert not oob, f"shapes off-canvas: {oob[:4]}"

    @pytest.mark.parametrize("density", ["airy", "normal", "dense"])
    def test_densities(self, density, tmp_path):
        deck = _full_deck(density=density)
        overflow, oob = _measure_pptx_overflow(
            render_deck(deck, str(tmp_path / "d.pptx"))
        )
        assert not overflow and not oob

    def test_four_three_aspect(self, tmp_path):
        from pptx import Presentation

        out = render_deck(_full_deck(aspect="4:3"), str(tmp_path / "a.pptx"))
        prs = Presentation(out)
        assert prs.slide_width / EMU_PER_PT == pytest.approx(720.0)
        overflow, oob = _measure_pptx_overflow(out)
        assert not overflow and not oob

    def test_slide_count_and_notes(self, tmp_path):
        from pptx import Presentation

        deck = _full_deck()
        prs = Presentation(render_deck(deck, str(tmp_path / "n.pptx")))
        assert len(prs.slides) == deck.slide_count()
        notes = [s.notes_slide.notes_text_frame.text for s in prs.slides
                 if s.has_notes_slide and s.notes_slide.notes_text_frame.text]
        assert "Speaker notes here." in notes

    def test_charts_and_tables_are_native(self, tmp_path):
        from pptx import Presentation

        prs = Presentation(render_deck(_full_deck(), str(tmp_path / "c.pptx")))
        charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        tables = sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)
        assert charts == 2, "charts must be native, editable chart parts"
        assert tables == 1, "tables must be native table parts"

    def test_creates_missing_directories(self, tmp_path):
        out = render_deck(_full_deck(), str(tmp_path / "deep" / "nested" / "d.pptx"))
        assert os.path.exists(out)

    def test_empty_and_degenerate_slides(self, tmp_path):
        """Missing payloads must render a slide, not raise."""
        deck = Deck(title="Edge", slides=[
            Slide(type=SlideType.CHART, title="No data"),
            Slide(type=SlideType.TABLE, title="No table"),
            Slide(type=SlideType.STATS, title="No stats"),
            Slide(type=SlideType.CARDS, title="No cards"),
            Slide(type=SlideType.TIMELINE, title="No items"),
            Slide(type=SlideType.PROCESS, title="No steps"),
            Slide(type=SlideType.COMPARISON, title="No sides"),
            Slide(type=SlideType.BULLETS, title=""),
            Slide(type=SlideType.IMAGE, title="Missing", image_path="/nope/none.png"),
        ])
        assert os.path.exists(render_deck(deck, str(tmp_path / "e.pptx")))

    def test_very_long_strings_do_not_overflow(self, tmp_path):
        deck = Deck(title="Long " * 40, slides=[
            Slide(type=SlideType.COVER, title="A title that simply refuses to stop " * 6),
            Slide(type=SlideType.BIG_STATEMENT, title="Statement " * 40),
            Slide(type=SlideType.BULLETS, title="Wordy " * 20,
                  bullets=[Bullet(text="A bullet " * 60)]),
            Slide(type=SlideType.CARDS, title="Cards", cards=[
                Card(title="Supercalifragilisticexpialidocious", body="Body " * 40)]),
        ])
        overflow, oob = _measure_pptx_overflow(render_deck(deck, str(tmp_path / "l.pptx")))
        assert not overflow, f"{overflow[:3]}"
        assert not oob


class TestPdfRenderer:
    @pytest.mark.parametrize("theme", ["kimi", "aurora", "editorial", "noir", "slate"])
    def test_renders_valid_pdf(self, theme, tmp_path):
        out = render_pdf(_full_deck(theme), str(tmp_path / f"{theme}.pdf"))
        assert os.path.getsize(out) > 5000
        with open(out, "rb") as fh:
            assert fh.read(5) == b"%PDF-"

    def test_page_count_matches_slides(self, tmp_path):
        fitz = pytest.importorskip("pymupdf")
        deck = _full_deck()
        out = render_pdf(deck, str(tmp_path / "p.pdf"))
        with fitz.open(out) as doc:
            assert len(doc) == deck.slide_count()

    def test_text_is_extractable_and_bookmarked(self, tmp_path):
        """A PDF should be searchable and navigable, not a picture of a deck."""
        fitz = pytest.importorskip("pymupdf")
        out = render_pdf(_full_deck(), str(tmp_path / "t.pdf"))
        with fitz.open(out) as doc:
            text = "\n".join(page.get_text() for page in doc)
            assert "One claim worth a slide." in text
            assert "A memorable line about the work." in text
            assert len(doc.get_toc()) == _full_deck().slide_count()

    def test_chart_types_all_draw(self, tmp_path):
        for kind in ("column", "bar", "line", "area", "pie", "doughnut",
                     "stacked_column", "stacked_bar"):
            deck = Deck(title=kind, slides=[Slide(type=SlideType.CHART, title=kind, chart=ChartData(
                chart_type=kind, categories=["a", "b", "c"], data_labels=True,
                series=[ChartSeries("S1", [3, 7, 5]), ChartSeries("S2", [6, 2, 8])]))])
            assert os.path.exists(render_pdf(deck, str(tmp_path / f"{kind}.pdf")))

    def test_non_latin_text_survives(self, tmp_path):
        """CJK must not crash the PDF path even when core fonts cannot encode it."""
        deck = Deck(title="渲染测试", slides=[
            Slide(type=SlideType.COVER, title="可再生能源现状报告", subtitle="2026 年展望"),
            Slide(type=SlideType.BULLETS, title="要点", bullets=[Bullet(text="成本持续下降")]),
        ])
        out = render_pdf(deck, str(tmp_path / "cjk.pdf"))
        assert os.path.getsize(out) > 1000

    def test_engine_soffice_errors_clearly_when_absent(self, tmp_path):
        from kimiagent.pdf_renderer import soffice_available

        if soffice_available():
            pytest.skip("LibreOffice is installed; nothing to assert")
        with pytest.raises(RuntimeError, match="LibreOffice"):
            render_pdf(_full_deck(), str(tmp_path / "s.pdf"), engine="soffice")

    def test_engine_auto_falls_back(self, tmp_path):
        assert os.path.exists(
            render_pdf(_full_deck(), str(tmp_path / "a.pdf"), engine="auto")
        )


# --------------------------------------------------------------------------
# agent.py
# --------------------------------------------------------------------------


class TestAgent:
    def test_offline_needs_no_key(self):
        deck = DeckAgent(api_key=None).build_offline("A topic about nothing")
        assert deck.slide_count() >= 10
        assert deck.slides[0].type is SlideType.COVER
        assert deck.slides[-1].type is SlideType.CLOSING

    def test_offline_output_is_varied(self):
        """The whole point: a deck must not be a stack of bullet slides."""
        deck = DeckAgent(api_key=None).build_offline("Renewable energy", slide_count=20)
        assert deck.variety_score() >= 0.5
        kinds = set(deck.type_histogram())
        assert {"big_statement", "stats", "cards"} <= kinds

    @pytest.mark.parametrize("budget", [8, 12, 16, 20, 26])
    def test_respects_slide_budget(self, budget):
        deck = DeckAgent(api_key=None).build_offline("Topic", slide_count=budget)
        assert abs(deck.slide_count() - budget) <= 4

    def test_explicit_outline_is_never_dropped(self):
        outline = "Highlights; Revenue; Product; Risks; Next quarter"
        deck = DeckAgent(api_key=None).build_offline("Q3", extra=outline, slide_count=10)
        titles = [s.title for s in deck.slides if s.type is SlideType.SECTION]
        assert titles == ["Highlights", "Revenue", "Product", "Risks", "Next quarter"]

    def test_audience_appears_on_cover(self):
        deck = DeckAgent(api_key=None).build_offline("T", audience="Board of directors")
        assert "Board of directors" in deck.slides[0].subtitle
        assert deck.author == ""

    def test_no_mangled_casing_in_generated_prose(self):
        """Regression: title-cased topics were lowercased word-by-word."""
        deck = DeckAgent(api_key=None).build_offline("the state of renewable energy in 2026")
        prose = " ".join([s.title for s in deck.slides] +
                         [s.subtitle for s in deck.slides] +
                         [s.quote for s in deck.slides])
        assert "renewable Energy" not in prose

    def test_generated_decks_render_both_formats(self, tmp_path):
        deck = DeckAgent(api_key=None).build_offline("Renewables", slide_count=20, theme="glacier")
        overflow, oob = _measure_pptx_overflow(
            render_deck(deck, str(tmp_path / "a.pptx"))
        )
        assert not overflow and not oob
        assert os.path.exists(render_pdf(deck, str(tmp_path / "a.pdf")))

    def test_critique_flags_a_bullet_only_deck(self):
        deck = Deck(title="x", slides=[
            Slide(type=SlideType.BULLETS, title=f"S{i}", bullets=[Bullet(text="p")])
            for i in range(6)
        ])
        notes = " ".join(critique(deck))
        assert "variety" in notes.lower()

    def test_critique_flags_chart_without_takeaway(self):
        deck = Deck(title="x", slides=[Slide(type=SlideType.CHART, title="c", chart=ChartData(
            chart_type="line", categories=["a"], series=[ChartSeries("s", [1])]))])
        assert any("takeaway" in n for n in critique(deck))

    def test_llm_failure_degrades_to_offline(self, monkeypatch):
        agent = DeckAgent(api_key="sk-not-real")

        def boom(*a, **k):
            raise RuntimeError("network down")

        monkeypatch.setattr(agent, "_build_with_llm", boom)
        deck = agent.build("A topic", use_llm=True)
        assert deck.slide_count() > 0

    def test_ensure_structure_repairs_llm_output(self):
        from kimiagent.agent import _ensure_structure

        deck = Deck(title="T", slides=[Slide(type=SlideType.BULLETS, title="Mid")])
        fixed = _ensure_structure(deck, "topic", "")
        assert fixed.slides[0].type is SlideType.COVER
        assert fixed.slides[-1].type is SlideType.CLOSING


# --------------------------------------------------------------------------
# cli.py
# --------------------------------------------------------------------------


class TestCli:
    def test_generate_both_formats(self, tmp_path):
        from kimiagent.cli import main

        stem = tmp_path / "out"
        assert main(["generate", "-t", "CLI test", "-f", "both", "-o", str(stem),
                     "--report", "--offline"]) == 0
        assert (tmp_path / "out.pptx").exists()
        assert (tmp_path / "out.pdf").exists()

    def test_explicit_extension_wins(self, tmp_path):
        from kimiagent.cli import main

        target = tmp_path / "named.pdf"
        assert main(["generate", "-t", "T", "-f", "pdf", "-o", str(target), "--offline"]) == 0
        assert target.exists()

    def test_emit_spec_round_trips_through_from_spec(self, tmp_path):
        from kimiagent.cli import main

        spec = tmp_path / "spec.json"
        assert main(["generate", "-t", "Round trip", "--offline", "--emit-spec", str(spec),
                     "-o", str(tmp_path / "a.pptx")]) == 0
        assert main(["from-spec", str(spec), "-f", "both",
                     "-o", str(tmp_path / "b"), "--theme", "aurora"]) == 0
        assert (tmp_path / "b.pdf").exists()
        assert Deck.from_dict(json.loads(spec.read_text())).slide_count() > 0

    def test_critique_and_themes_commands(self, tmp_path, capsys):
        from kimiagent.cli import main

        spec = tmp_path / "s.json"
        spec.write_text(json.dumps(_full_deck().to_dict()))
        assert main(["critique", str(spec)]) == 0
        assert main(["themes"]) == 0
        assert "kimi" in capsys.readouterr().out

    def test_missing_file_exits_nonzero(self, tmp_path):
        from kimiagent.cli import main

        assert main(["from-spec", str(tmp_path / "nope.json")]) == 1

    def test_invalid_json_exits_nonzero(self, tmp_path):
        from kimiagent.cli import main

        bad = tmp_path / "bad.json"
        bad.write_text("{not json")
        assert main(["from-spec", str(bad)]) == 1
