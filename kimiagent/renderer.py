"""Render a :class:`~kimiagent.models.Deck` into a native, *modern* ``.pptx``.

This renderer aims for a contemporary, editorial "deck design" aesthetic —
gradient canvases, rounded cards with soft shadows, oversized display type,
ghosted section numerals, pill/chip labels and accent geometry — while still
emitting fully native, editable PowerPoint (real shapes, native charts and
native tables). Only ``python-pptx`` is required; Pillow is used opportunistically
for image sizing.
"""

from __future__ import annotations

from typing import List, Optional, Tuple
import struct

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

from .models import Deck, Slide, SlideType, Bullet, Stat, ChartData, TableData, ImageItem
from .themes import Theme, get_theme


# 16:9 canvas.
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.9)
CONTENT_W = SLIDE_W - 2 * MARGIN


# --------------------------------------------------------------------------
# Colour helpers
# --------------------------------------------------------------------------

def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _luminance(hex_str: str) -> float:
    r = int(hex_str[0:2], 16) / 255.0
    g = int(hex_str[2:4], 16) / 255.0
    b = int(hex_str[4:6], 16) / 255.0
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _contrast(a: str, b: str) -> float:
    la, lb = _luminance(a), _luminance(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def _shift(hex_str: str, amt: float) -> str:
    """Lighten (amt>0) or darken (amt<0) a colour toward white/black."""
    r = int(hex_str[0:2], 16)
    g = int(hex_str[2:4], 16)
    b = int(hex_str[4:6], 16)
    if amt >= 0:
        r = int(r + (255 - r) * amt)
        g = int(g + (255 - g) * amt)
        b = int(b + (255 - b) * amt)
    else:
        f = 1 + amt
        r, g, b = int(r * f), int(g * f), int(b * f)
    return f"{max(0,min(255,r)):02X}{max(0,min(255,g)):02X}{max(0,min(255,b)):02X}"


class DeckRenderer:
    """Renders a deck with a modern visual system."""

    def __init__(self, deck: Deck) -> None:
        self.deck = deck
        self.theme: Theme = get_theme(deck.theme)
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self._blank = self.prs.slide_layouts[6]
        self._section_no = 0
        self._dark = _luminance(self.theme.background) < 0.5

    # -- entry point ------------------------------------------------------

    def render(self) -> Presentation:
        dispatch = {
            SlideType.COVER: self._cover,
            SlideType.AGENDA: self._agenda,
            SlideType.SECTION: self._section,
            SlideType.BULLETS: self._bullets,
            SlideType.TWO_COLUMN: self._two_column,
            SlideType.STATS: self._stats,
            SlideType.CHART: self._chart,
            SlideType.TABLE: self._table,
            SlideType.QUOTE: self._quote,
            SlideType.IMAGE: self._image,
            SlideType.GALLERY: self._gallery,
            SlideType.CLOSING: self._closing,
        }
        total = len(self.deck.slides)
        for idx, slide in enumerate(self.deck.slides):
            builder = dispatch.get(slide.type, self._bullets)
            pptx_slide = self.prs.slides.add_slide(self._blank)
            builder(pptx_slide, slide)
            self._add_notes(pptx_slide, slide.notes)
            if slide.type not in (SlideType.COVER, SlideType.SECTION, SlideType.CLOSING, SlideType.QUOTE):
                self._add_footer(pptx_slide, idx + 1, total)
        return self.prs

    # ------------------------------------------------------------------
    # Low-level primitives
    # ------------------------------------------------------------------

    def _bg_solid(self, slide, hex_color: str) -> None:
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = _rgb(hex_color)

    def _bg_gradient(self, slide, c1: str, c2: str, angle: float = 60.0) -> None:
        """Full-bleed linear-gradient canvas."""
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        shape.line.fill.background()
        shape.shadow.inherit = False
        try:
            shape.fill.gradient()
            stops = shape.fill.gradient_stops
            stops[0].position = 0.0
            stops[0].color.rgb = _rgb(c1)
            stops[-1].position = 1.0
            stops[-1].color.rgb = _rgb(c2)
            try:
                shape.fill.gradient_angle = angle
            except Exception:
                pass
        except Exception:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(c1)
        # send to back
        sp = shape._element
        sp.getparent().remove(sp)
        slide.shapes._spTree.insert(2, sp)
        return shape

    def _add_shadow(self, shape, blur=90000, dist=45000, direction=5400000, alpha=26000) -> None:
        try:
            spPr = shape._element.spPr
            if spPr.find(qn("a:effectLst")) is not None:
                return
            xml = (
                f'<a:effectLst {nsdecls("a")}>'
                f'<a:outerShdw blurRad="{blur}" dist="{dist}" dir="{direction}" rotWithShape="0">'
                f'<a:srgbClr val="000000"><a:alpha val="{alpha}"/></a:srgbClr>'
                f'</a:outerShdw></a:effectLst>'
            )
            spPr.append(parse_xml(xml))
        except Exception:
            pass

    def _round(self, shape, radius: float = 0.09) -> None:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass

    def _rect(self, slide, left, top, width, height, hex_color, shape_type=MSO_SHAPE.RECTANGLE,
              line_color: Optional[str] = None, shadow: bool = False, radius: Optional[float] = None):
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(hex_color)
        if line_color:
            shape.line.color.rgb = _rgb(line_color)
            shape.line.width = Pt(1.0)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            self._round(shape, radius)
        if shadow:
            self._add_shadow(shape)
        return shape

    def _card(self, slide, left, top, width, height, fill=None, radius=0.06, shadow=True):
        fill = fill or self.theme.surface
        return self._rect(slide, left, top, width, height, fill,
                          shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=shadow, radius=radius)

    def _oval(self, slide, left, top, width, height, hex_color):
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(hex_color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def _text(self, slide, left, top, width, height, text, size, color, bold=False,
              italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=None,
              line_spacing=None, spacing=None):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = _rgb(color)
        run.font.name = font or self.theme.body_font
        if spacing is not None:
            self._letter_spacing(run, spacing)
        return box

    def _letter_spacing(self, run, pts: float) -> None:
        try:
            run._r.get_or_add_rPr().set("spc", str(int(pts * 100)))
        except Exception:
            pass

    def _chip(self, slide, left, top, text, fill, text_color, size=11, bold=True):
        w = Inches(0.42 + 0.098 * len(text) * (size / 11.0))
        h = Inches(0.06 * size + 0.06)
        shape = self._rect(slide, left, top, w, h, fill,
                           shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tf = shape.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = _rgb(text_color)
        r.font.name = self.theme.body_font
        self._letter_spacing(r, 0.6)
        return shape, w, h

    def _dots(self, slide, left, top, color, n=3, d=Inches(0.12), gap=Inches(0.11)):
        for i in range(n):
            self._oval(slide, left + i * (d + gap), top, d, d, color)

    def _ghost_number(self, slide, text, left, top, size=320, color=None):
        color = color or (_shift(self.theme.background, 0.10) if self._dark else _shift(self.theme.text, 0.86))
        self._text(slide, left, top, Inches(7), Inches(5.2), text, size, color,
                   bold=True, font=self.theme.heading_font)

    # -- shared heading motif --------------------------------------------

    def _title_color(self, bg_hex: str) -> str:
        if _contrast(self.theme.primary, bg_hex) >= 2.8:
            return self.theme.primary
        return self.theme.text

    def _kicker(self, slide, text, left=None, top=None):
        left = MARGIN if left is None else left
        top = Inches(0.62) if top is None else top
        self._chip(slide, left, top, text.upper(), self.theme.secondary, self.theme.text_on_primary, size=10)

    def _heading(self, slide, title: str, subtitle: str = "", kicker: str = ""):
        top = Inches(0.62)
        if kicker:
            self._kicker(slide, kicker, MARGIN, top)
            top = Inches(1.18)
        # Shrink long titles so they stay on one clean line.
        size = 34
        if len(title) > 44:
            size = 26
        elif len(title) > 32:
            size = 30
        self._text(slide, MARGIN, top, CONTENT_W, Inches(0.95), title, size,
                   self._title_color(self.theme.background), bold=True, font=self.theme.heading_font)
        # accent underline
        self._rect(slide, MARGIN, top + Inches(0.84), Inches(0.9), Inches(0.08), self.theme.secondary)
        if subtitle:
            self._text(slide, MARGIN + Inches(1.05), top + Inches(0.74), CONTENT_W - Inches(1.05),
                       Inches(0.4), subtitle, 14, self.theme.text_muted, italic=True)

    def _content_top(self, has_kicker=True):
        return Inches(2.05) if has_kicker else Inches(1.7)

    def _add_footer(self, slide, page: int, total: int):
        self._oval(slide, MARGIN, SLIDE_H - Inches(0.52), Inches(0.1), Inches(0.1), self.theme.secondary)
        self._text(slide, MARGIN + Inches(0.2), SLIDE_H - Inches(0.62), Inches(8), Inches(0.3),
                   self.deck.title, 9, self.theme.text_muted)
        self._text(slide, SLIDE_W - MARGIN - Inches(2), SLIDE_H - Inches(0.62), Inches(2), Inches(0.3),
                   f"{page:02d} / {total:02d}", 9, self.theme.text_muted, align=PP_ALIGN.RIGHT)

    def _add_notes(self, slide, notes: str):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # ------------------------------------------------------------------
    # Slide builders
    # ------------------------------------------------------------------

    def _cover(self, slide, s: Slide):
        p = self.theme.primary
        self._bg_gradient(slide, _shift(p, -0.25), _shift(p, 0.12), angle=55)
        # accent geometry: large circle bleeding off the top-right
        self._oval(slide, SLIDE_W - Inches(3.1), -Inches(2.2), Inches(5.2), Inches(5.2), self.theme.secondary)
        self._oval(slide, SLIDE_W - Inches(1.2), Inches(2.3), Inches(1.5), Inches(1.5),
                   _shift(self.theme.secondary, 0.25))
        # brand kicker
        self._chip(slide, MARGIN, Inches(0.95), "DIGITAL MARKETING AGENCY · GOA",
                   _shift(p, 0.16), self.theme.text_on_primary, size=11)
        # big title
        self._text(slide, MARGIN, Inches(2.7), Inches(9.6), Inches(2.6),
                   s.title or self.deck.title, 60, self.theme.text_on_primary, bold=True,
                   font=self.theme.heading_font, line_spacing=1.0)
        # accent bar
        self._rect(slide, MARGIN + Inches(0.03), Inches(4.55), Inches(1.4), Inches(0.11), self.theme.secondary)
        if s.subtitle:
            self._text(slide, MARGIN, Inches(4.8), Inches(9.2), Inches(1.0), s.subtitle, 20,
                       _shift(self.theme.text_on_primary, -0.06), italic=False, line_spacing=1.1)
        self._dots(slide, MARGIN, SLIDE_H - Inches(1.05), self.theme.secondary, n=3)
        if self.deck.author:
            self._text(slide, SLIDE_W - MARGIN - Inches(4), SLIDE_H - Inches(1.1), Inches(4), Inches(0.4),
                       self.deck.author, 13, _shift(self.theme.text_on_primary, -0.1), align=PP_ALIGN.RIGHT)

    def _closing(self, slide, s: Slide):
        p = self.theme.primary
        self._bg_gradient(slide, _shift(p, -0.25), _shift(p, 0.14), angle=55)
        self._oval(slide, -Inches(1.8), SLIDE_H - Inches(3.0), Inches(4.6), Inches(4.6),
                   _shift(self.theme.secondary, 0.0))
        self._chip(slide, MARGIN, Inches(2.0), "GET IN TOUCH", _shift(p, 0.16),
                   self.theme.text_on_primary, size=11)
        self._text(slide, MARGIN, Inches(2.7), Inches(11.0), Inches(1.7),
                   s.title or "Let's work together", 54, self.theme.text_on_primary, bold=True,
                   font=self.theme.heading_font, line_spacing=1.0)
        self._rect(slide, MARGIN + Inches(0.03), Inches(4.4), Inches(1.4), Inches(0.11), self.theme.secondary)
        if s.subtitle:
            self._text(slide, MARGIN, Inches(4.7), Inches(11.0), Inches(0.9), s.subtitle, 18,
                       _shift(self.theme.text_on_primary, -0.05))
        self._dots(slide, MARGIN, SLIDE_H - Inches(1.0), self.theme.secondary, n=3)

    def _section(self, slide, s: Slide):
        self._section_no += 1
        p = self.theme.primary
        self._bg_gradient(slide, _shift(p, -0.18), _shift(p, 0.16), angle=50)
        # ghosted numeral
        self._ghost_number(slide, f"{self._section_no:02d}", SLIDE_W - Inches(6.4), Inches(0.2),
                           size=340, color=_shift(p, 0.1))
        self._chip(slide, MARGIN, Inches(2.55), (s.subtitle or "SECTION").upper(),
                   self.theme.secondary, self.theme.text_on_primary, size=11)
        self._text(slide, MARGIN, Inches(3.15), Inches(9.5), Inches(1.8), s.title, 46,
                   self.theme.text_on_primary, bold=True, font=self.theme.heading_font, line_spacing=1.0)
        self._rect(slide, MARGIN + Inches(0.03), Inches(4.85), Inches(1.4), Inches(0.11), self.theme.secondary)

    def _bullets(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        # slim corner accent (top-right) for rhythm without crowding text
        self._rect(slide, SLIDE_W - Inches(1.9), 0, Inches(1.9), Inches(0.16), self.theme.secondary)
        self._heading(slide, s.title, s.subtitle)
        top = Inches(2.15)
        self._modern_bullets(slide, s.bullets, MARGIN, top, CONTENT_W - Inches(0.3),
                             SLIDE_H - top - Inches(0.9))

    def _modern_bullets(self, slide, bullets: List[Bullet], left, top, width, height):
        if not bullets:
            return
        n = len(bullets)
        row_h = min(Inches(1.0), height / max(n, 1))
        for i, b in enumerate(bullets):
            y = top + i * row_h
            indent = Inches(0.5) * b.level
            # marker
            m = Inches(0.26)
            self._rect(slide, left + indent, y + Inches(0.08), m, m,
                       self.theme.accent(i), shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            tx = left + indent + m + Inches(0.3)
            tw = width - indent - m - Inches(0.3)
            box = slide.shapes.add_textbox(tx, y - Inches(0.02), tw, row_h)
            tf = box.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.line_spacing = 1.05
            size = 20 if b.level == 0 else 16
            if b.bold_lead:
                lead = para.add_run()
                lead.text = f"{b.bold_lead}  "
                lead.font.bold = True
                lead.font.size = Pt(size)
                lead.font.color.rgb = _rgb(self.theme.secondary)
                lead.font.name = self.theme.body_font
            rest = para.add_run()
            rest.text = b.text
            rest.font.size = Pt(size)
            rest.font.bold = b.level == 0 and not b.bold_lead
            rest.font.color.rgb = _rgb(self.theme.text)
            rest.font.name = self.theme.body_font

    def _agenda(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title or "Agenda", kicker="OVERVIEW")
        items = s.bullets
        top = self._content_top()
        avail = SLIDE_H - top - Inches(0.8)
        cols = 2 if len(items) > 4 else 1
        rows = (len(items) + cols - 1) // cols
        col_w = (CONTENT_W - Inches(0.5)) / cols
        row_h = min(Inches(1.15), avail / max(rows, 1))
        for i, b in enumerate(items):
            r = i % rows
            c = i // rows
            x = MARGIN + c * (col_w + Inches(0.5))
            y = top + r * row_h
            self._text(slide, x, y - Inches(0.05), Inches(1.1), Inches(0.9), f"{i+1:02d}",
                       30, self.theme.secondary, bold=True, font=self.theme.heading_font)
            self._text(slide, x + Inches(1.05), y, col_w - Inches(1.05), Inches(0.8), b.text,
                       18, self.theme.text, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            self._rect(slide, x + Inches(1.05), y + row_h - Inches(0.22), col_w - Inches(1.25),
                       Emu(9525), _shift(self.theme.text_muted, 0.4))

    def _two_column(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title, s.subtitle, kicker="COMPARE")
        cols = s.columns[:2] if s.columns else [[], []]
        col_w = (CONTENT_W - Inches(0.5)) / 2
        top = self._content_top()
        card_h = SLIDE_H - top - Inches(0.8)
        for ci, col in enumerate(cols):
            left = MARGIN + ci * (col_w + Inches(0.5))
            card = self._card(slide, left, top, col_w, card_h, fill=self.theme.surface, radius=0.05)
            self._rect(slide, left, top, col_w, Inches(0.16), self.theme.accent(ci),
                       shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            self._modern_bullets(slide, col, left + Inches(0.4), top + Inches(0.55),
                                 col_w - Inches(0.8), card_h - Inches(0.9))

    def _stats(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title, s.subtitle, kicker="BY THE NUMBERS")
        stats = s.stats[:4]
        if not stats:
            return
        n = len(stats)
        gap = Inches(0.4)
        card_w = (CONTENT_W - gap * (n - 1)) / n
        top = Inches(2.55)
        card_h = Inches(3.1)
        for i, st in enumerate(stats):
            left = MARGIN + i * (card_w + gap)
            self._card(slide, left, top, card_w, card_h, fill=self.theme.surface, radius=0.07)
            self._rect(slide, left + Inches(0.35), top + Inches(0.4), Inches(0.5), Inches(0.1),
                       self.theme.accent(i))
            self._text(slide, left + Inches(0.3), top + Inches(0.65), card_w - Inches(0.6), Inches(1.3),
                       st.value, 52, self.theme.accent(i), bold=True, font=self.theme.heading_font)
            self._text(slide, left + Inches(0.32), top + Inches(1.95), card_w - Inches(0.6), Inches(0.6),
                       st.label, 15, self.theme.text, bold=True)
            if st.detail:
                self._text(slide, left + Inches(0.32), top + Inches(2.45), card_w - Inches(0.6),
                           Inches(0.55), st.detail, 11, self.theme.text_muted)

    def _chart(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title, s.subtitle, kicker="DATA")
        cd = s.chart
        top = self._content_top()
        # white card behind chart for legibility on any theme
        card = self._card(slide, MARGIN, top, CONTENT_W, SLIDE_H - top - Inches(0.8),
                          fill="FFFFFF", radius=0.04)
        if not cd or not cd.series:
            self._text(slide, MARGIN, top + Inches(1), CONTENT_W, Inches(1), "(No chart data)",
                       16, self.theme.text_muted, align=PP_ALIGN.CENTER)
            return
        chart_data = CategoryChartData()
        chart_data.categories = cd.categories or [str(i + 1) for i in range(len(cd.series[0].values))]
        for series in cd.series:
            chart_data.add_series(series.name, tuple(series.values))
        x, y = MARGIN + Inches(0.3), top + Inches(0.3)
        cx, cy = CONTENT_W - Inches(0.6), SLIDE_H - top - Inches(1.4)
        gframe = slide.shapes.add_chart(_chart_type(cd.chart_type), x, y, cx, cy, chart_data)
        chart = gframe.chart
        chart.has_title = False
        multi = len(cd.series) > 1 or cd.chart_type == "pie"
        chart.has_legend = multi
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(11)
        try:
            self._color_chart(chart, cd)
        except Exception:
            pass

    def _color_chart(self, chart, cd: ChartData):
        plot = chart.plots[0]
        if cd.chart_type == "pie":
            series = plot.series[0]
            for i, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(self.theme.accent(i))
        else:
            for i, series in enumerate(plot.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = _rgb(self.theme.accent(i))
                series.format.line.color.rgb = _rgb(self.theme.accent(i))

    def _table(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title, s.subtitle, kicker="DETAIL")
        td = s.table
        if not td or not td.headers:
            return
        rows = len(td.rows) + 1
        cols = len(td.headers)
        top = self._content_top()
        height = min(SLIDE_H - top - Inches(0.9), Inches(0.62) * rows)
        gframe = slide.shapes.add_table(rows, cols, MARGIN, top, CONTENT_W, height)
        table = gframe.table
        for c, header in enumerate(td.headers):
            cell = table.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(self.theme.primary if not self._dark else self.theme.secondary)
            self._style_cell(cell, header, bold=True, color=self.theme.text_on_primary, size=14)
        for r, row in enumerate(td.rows, start=1):
            band = self.theme.surface if r % 2 == 1 else _shift(self.theme.surface, 0.06 if self._dark else -0.02)
            for c in range(cols):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(band)
                value = row[c] if c < len(row) else ""
                self._style_cell(cell, value, bold=False, color=self.theme.text, size=13)

    def _style_cell(self, cell, text, bold, color, size):
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.2)
        cell.margin_right = Inches(0.15)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = str(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        run.font.name = self.theme.body_font

    def _quote(self, slide, s: Slide):
        p = self.theme.primary
        self._bg_gradient(slide, _shift(p, -0.15), _shift(p, 0.2), angle=45)
        self._text(slide, MARGIN - Inches(0.1), Inches(0.7), Inches(3), Inches(2.4), "\u201C", 200,
                   self.theme.secondary, bold=True, font=self.theme.heading_font)
        self._text(slide, Inches(1.7), Inches(2.25), Inches(10.2), Inches(3.0), s.quote, 27,
                   self.theme.text_on_primary, bold=True, anchor=MSO_ANCHOR.MIDDLE,
                   font=self.theme.heading_font, line_spacing=1.16)
        if s.attribution:
            self._chip(slide, Inches(1.75), Inches(5.85), s.attribution, self.theme.secondary,
                       self.theme.text_on_primary, size=13)

    def _image(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        if s.title:
            self._heading(slide, s.title, s.subtitle, kicker="SHOWCASE")
        top = self._content_top() if s.title else Inches(0.6)
        bottom = Inches(0.85)
        box_h = SLIDE_H - top - bottom
        card = self._card(slide, MARGIN, top, CONTENT_W, box_h, fill=_shift(self.theme.surface, 0.02), radius=0.04)
        cap_h = Inches(0.5) if s.caption else Inches(0.0)
        placed = self._place_image_contained(slide, s.image_path,
                                              MARGIN + Inches(0.2), top + Inches(0.2),
                                              CONTENT_W - Inches(0.4), box_h - Inches(0.4) - cap_h,
                                              rounded=True, shadow=True)
        if not placed:
            self._text(slide, MARGIN, top + box_h / 2 - Inches(0.4), CONTENT_W, Inches(0.8),
                       s.caption or "[ image ]", 16, self.theme.text_muted,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if s.caption:
            self._text(slide, MARGIN, top + box_h - Inches(0.5), CONTENT_W, Inches(0.4),
                       s.caption, 13, self.theme.text_muted, italic=True, align=PP_ALIGN.CENTER)

    def _gallery(self, slide, s: Slide):
        self._bg_solid(slide, self.theme.background)
        self._heading(slide, s.title, s.subtitle, kicker="PORTFOLIO")
        items = s.images
        if not items:
            return
        n = len(items)
        cols = _grid_columns(n)
        rows = (n + cols - 1) // cols
        area_left = MARGIN
        area_top = self._content_top()
        area_w = CONTENT_W
        area_h = SLIDE_H - area_top - Inches(0.75)
        gutter = Inches(0.34)
        cell_w = (area_w - gutter * (cols - 1)) / cols
        cell_h = (area_h - gutter * (rows - 1)) / rows
        cap_h = Inches(0.4)
        for i, item in enumerate(items):
            r, c = divmod(i, cols)
            cx = area_left + c * (cell_w + gutter)
            cy = area_top + r * (cell_h + gutter)
            # white mat card with shadow
            self._card(slide, cx, cy, cell_w, cell_h, fill="FFFFFF", radius=0.05)
            img_h = cell_h - (cap_h if item.caption else Inches(0.15))
            placed = self._place_image_contained(
                slide, item.path, cx + Inches(0.12), cy + Inches(0.12),
                cell_w - Inches(0.24), img_h - Inches(0.16), rounded=True, shadow=False)
            if not placed:
                self._text(slide, cx, cy + img_h / 2 - Inches(0.3), cell_w, Inches(0.6),
                           "[ image ]", 12, self.theme.text_muted, align=PP_ALIGN.CENTER,
                           anchor=MSO_ANCHOR.MIDDLE)
            if item.caption:
                self._text(slide, cx + Inches(0.12), cy + img_h - Inches(0.04), cell_w - Inches(0.24),
                           cap_h, item.caption, 10.5, "3A3A3A", align=PP_ALIGN.CENTER,
                           anchor=MSO_ANCHOR.MIDDLE, bold=True)

    # -- image placement --------------------------------------------------

    def _place_image_contained(self, slide, path, box_left, box_top, box_w, box_h,
                               rounded=False, shadow=False):
        if not path:
            return False
        size = _image_size(path)
        if size is None:
            try:
                pic = slide.shapes.add_picture(path, box_left, box_top, width=box_w)
                return True
            except Exception:
                return False
        iw, ih = size
        if iw <= 0 or ih <= 0:
            return False
        scale = min(box_w / iw, box_h / ih)
        w = int(iw * scale)
        h = int(ih * scale)
        left = int(box_left + (box_w - w) / 2)
        top = int(box_top + (box_h - h) / 2)
        try:
            pic = slide.shapes.add_picture(path, left, top, width=w, height=h)
        except Exception:
            return False
        if rounded:
            self._round_picture(pic, 0.045)
        if shadow:
            self._add_shadow(pic, alpha=22000)
        return True

    def _round_picture(self, pic, val: float = 0.05):
        try:
            spPr = pic._element.spPr
            for tag in ("a:prstGeom", "a:custGeom"):
                existing = spPr.find(qn(tag))
                if existing is not None:
                    spPr.remove(existing)
            geom_xml = (
                f'<a:prstGeom {nsdecls("a")} prst="roundRect">'
                f'<a:avLst><a:gd name="adj" fmla="val {int(val*100000)}"/></a:avLst>'
                f'</a:prstGeom>'
            )
            geom = parse_xml(geom_xml)
            xfrm = spPr.find(qn("a:xfrm"))
            if xfrm is not None:
                xfrm.addnext(geom)
            else:
                spPr.append(geom)
        except Exception:
            pass


# --------------------------------------------------------------------------
# Module helpers
# --------------------------------------------------------------------------

def _grid_columns(n: int) -> int:
    return {1: 1, 2: 2, 3: 3, 4: 2, 5: 3, 6: 3}.get(n, 3 if n <= 9 else 4)


def _image_size(path):
    """Return (width, height) in pixels, or None. Pillow first, header fallback."""
    try:
        from PIL import Image  # type: ignore

        with Image.open(path) as im:
            return im.size
    except Exception:
        pass
    try:
        with open(path, "rb") as f:
            head = f.read(32)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            w, h = struct.unpack(">II", head[16:24])
            return int(w), int(h)
        if head[:2] == b"\xff\xd8":
            with open(path, "rb") as f:
                data = f.read()
            i = 2
            while i < len(data) - 9:
                if data[i] != 0xFF:
                    i += 1
                    continue
                marker = data[i + 1]
                if marker in (0xC0, 0xC1, 0xC2, 0xC3):
                    h = struct.unpack(">H", data[i + 5:i + 7])[0]
                    w = struct.unpack(">H", data[i + 7:i + 9])[0]
                    return int(w), int(h)
                seg_len = struct.unpack(">H", data[i + 2:i + 4])[0]
                i += 2 + seg_len
    except Exception:
        pass
    return None


def _chart_type(name: str) -> int:
    mapping = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "area": XL_CHART_TYPE.AREA,
    }
    return mapping.get((name or "column").lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)


def render_deck(deck: Deck, output_path: str) -> str:
    """Render ``deck`` to ``output_path`` (creating parent dirs) and return it."""
    import os

    parent = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(parent, exist_ok=True)
    renderer = DeckRenderer(deck)
    prs = renderer.render()
    prs.save(output_path)
    return output_path
